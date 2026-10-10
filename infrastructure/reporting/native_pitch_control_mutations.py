"""Single-target valid native mutations of actual configured Pitch bytes."""

from __future__ import annotations

from io import BytesIO
from typing import Any

from infrastructure.project.public_reference_git import remaining
from infrastructure.reporting.native_pitch_control_plan import PitchTarget

WRONG_TARGET = "https://example.invalid/native-pitch-control"


def _pdf(content: bytes, target: PitchTarget, *, deadline: float) -> bytes:
    from pypdf import PdfReader, PdfWriter
    from pypdf.generic import ContentStream, NameObject, TextStringObject
    from reportlab.pdfgen.canvas import Canvas

    document = PdfReader(BytesIO(content), strict=True)
    if target.page is None or not 1 <= target.page <= len(document.pages):
        raise ValueError("Pitch PDF control did not select one exact actual page")
    page = document.pages[target.page - 1]
    width, height = float(page.mediabox.width), float(page.mediabox.height)
    if target.name == "layout":
        page.mediabox.upper_right = (width - 12, height)
    elif target.name in {"citation", "qr-link"}:
        matches = []
        for reference in page.get("/Annots", []):
            remaining(deadline)
            annotation = reference.get_object()
            action = annotation.get("/A")
            if action is not None and str(action.get_object().get("/URI", "")) == target.source_target:
                matches.append(action.get_object())
        if len(matches) != 1:
            raise ValueError("Pitch PDF control requires one exact Source-selected native link")
        matches[0][NameObject("/URI")] = TextStringObject(WRONG_TARGET)
    elif target.name in {"theme", "qr-pixels"}:
        # merge_page establishes a balanced q/Q around both streams. The new
        # ReportLab stream sets its own CTM/colors, so existing page clipping or
        # transparency cannot silently turn this into an ineffective mutation.
        stream = BytesIO()
        canvas = Canvas(stream, pagesize=(width, height), invariant=1)
        if target.name == "theme":
            canvas.setFillColorRGB(1, 0, 1)
            canvas.rect(5, 5, 20, 20, fill=1, stroke=0)
        else:
            canvas.setFillColorRGB(1, 1, 1)
            # Actual Source Pitch QR geometry: top=height-53.28 in PPTX.
            canvas.rect(width - 39.6 - 44.64 - 2, 53.28 - 44.64 - 2, 48.64, 48.64, fill=1, stroke=0)
        canvas.save()
        overlay = PdfReader(BytesIO(stream.getvalue()), strict=True).pages[0]
        operations = ContentStream(overlay.get_contents(), overlay.pdf)
        operations.operations = [
            (args, op) for args, op in operations.operations if op not in {b"BT", b"Tf", b"TL", b"ET"}
        ]
        overlay[NameObject("/Contents")] = operations
        resources: Any = overlay["/Resources"]
        resources.pop("/Font", None)
        page.merge_page(overlay)
    else:
        raise ValueError("Pitch PDF cannot provide editable notes or an unknown native control")
    writer = PdfWriter()
    for item in document.pages:
        remaining(deadline)
        writer.add_page(item)
    output = BytesIO()
    writer.write(output)
    return output.getvalue()


def _pptx(content: bytes, target: PitchTarget, *, deadline: float) -> bytes:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    from pptx.util import Pt

    deck = Presentation(BytesIO(content))
    if target.page is None or not 1 <= target.page <= len(deck.slides):
        raise ValueError("Pitch PPTX control did not select one exact actual page")
    slide = deck.slides[target.page - 1]
    if target.name == "layout":
        shapes = [shape for shape in slide.shapes if shape.has_text_frame and shape.text_frame.text]
        if not shapes:
            raise ValueError("Pitch layout control lacks its actual Source text shape")
        shapes[0].left += Pt(12)
    elif target.name == "theme":
        # The fixed Source renderer makes the first shape its full-page native
        # background rectangle. Refuse rather than silently choosing another.
        shape = slide.shapes[0]
        if (shape.left, shape.top, shape.width, shape.height) != (0, 0, deck.slide_width, deck.slide_height):
            raise ValueError("Pitch theme control lacks the exact Source background geometry")
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(255, 0, 255)
    elif target.name == "notes":
        frame = slide.notes_slide.notes_text_frame
        if frame is None or frame.text != target.source_target:
            raise ValueError("Pitch note control differs from its exact Source-selected note")
        frame.text = target.source_target + "\nNative controlled note mutation."
    elif target.name in {"citation", "qr-link", "qr-pixels"}:
        links = []
        pictures = []
        for shape in slide.shapes:
            remaining(deadline)
            if (
                shape.shape_type == MSO_SHAPE_TYPE.PICTURE
                and shape.click_action.hyperlink.address == target.source_target
            ):
                pictures.append(shape)
                if target.name == "qr-link":
                    links.append(shape.click_action.hyperlink)
            if target.name == "citation" and shape.has_text_frame:
                links.extend(
                    run.hyperlink
                    for paragraph in shape.text_frame.paragraphs
                    for run in paragraph.runs
                    if run.hyperlink.address == target.source_target
                )
        if target.name == "qr-pixels":
            if len(pictures) != 1:
                raise ValueError("Pitch QR mutation requires one exact Source-targeted native picture")
            # A fresh opaque white picture keeps the same link/geometry and
            # avoids modifying any shared media part used by another page.
            from PIL import Image

            stream = BytesIO()
            with Image.new("RGB", (64, 64), "white") as image:
                image.save(stream, format="PNG")
            old = pictures[0]
            fresh = slide.shapes.add_picture(BytesIO(stream.getvalue()), old.left, old.top, old.width, old.height)
            fresh.click_action.hyperlink.address = target.source_target
            old._element.getparent().remove(old._element)
        else:
            if len(links) != 1:
                raise ValueError("Pitch control requires one exact Source-selected package link")
            links[0].address = WRONG_TARGET
    else:
        raise ValueError("unknown configured Pitch native package mutation")
    output = BytesIO()
    deck.save(output)
    return output.getvalue()


def mutate_pitch(content: bytes, kind: str, target: PitchTarget, *, deadline: float) -> bytes:
    """Leave every other actual page present; no invented-page fixture is used."""
    from infrastructure.reporting.native_format_acceptance_checks import parse_native_bytes

    remaining(deadline)
    if len(content) > 250 * 1024 * 1024:
        raise ValueError("Pitch control exceeds its native input byte budget")
    if kind == "pdf":
        from infrastructure.reporting.native_pdf_stream_bounds import preflight_pdf_streams

        preflight_pdf_streams(content, deadline=deadline, max_bytes=250 * 1024 * 1024, max_objects=2_000_000)
    parse_native_bytes(kind, content)
    if target.name == "baseline":
        if target.page is not None or target.intended_diagnostic is not None:
            raise ValueError("Pitch baseline cannot silently select a mutant page")
        return content
    result = (
        _pdf(content, target, deadline=deadline)
        if kind == "pdf"
        else _pptx(content, target, deadline=deadline)
        if kind == "pptx"
        else None
    )
    if result is None or result == content:
        raise ValueError("Pitch control must change actual selected native bytes")
    remaining(deadline)
    return result

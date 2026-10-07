"""Bounded native Office conversion; package parsing never grants layout acceptance."""

from __future__ import annotations

import shutil
from pathlib import Path
from typing import Any, TYPE_CHECKING

if TYPE_CHECKING:
    from infrastructure.reporting.native_format_content_contracts import NativeContentContract
    from infrastructure.reporting.native_projection_contracts import ProjectionPage

from infrastructure.core.files.secure_write import atomic_write_bytes_confined
from infrastructure.reporting.native_format_acceptance_checks import (
    NativeToolUnavailable,
    native_tool_record,
    parse_native_bytes,
)
from infrastructure.project.native_reader_runtime_custody import (
    observe_native_reader_runtime,
    retain_native_reader_runtime,
    validate_native_reader_runtime,
)
from infrastructure.reporting.release_readiness_snapshot import ArtifactSnapshot, SnapshotReader


def convert_office(
    root: Path,
    directory: Path,
    artifact: ArtifactSnapshot,
    kind: str,
    *,
    env: dict[str, str],
    deadline: float,
    run_id: str,
) -> dict[str, Any]:
    """Convert exact admitted DOCX/PPTX with an isolated LibreOffice profile.

    Caller supplies a fresh private evidence directory under root. Conversion
    is observed conformance only; source and visual checks remain mandatory.
    """
    if kind not in {"docx", "pptx"}:
        raise ValueError("native Office conversion requires DOCX or PPTX")
    selected = shutil.which("soffice", path=env.get("PATH"))
    if selected is None:
        raise NativeToolUnavailable("actual native LibreOffice reader is unavailable")
    try:
        runtime_before = observe_native_reader_runtime(Path(selected), deadline=deadline)
    except ValueError as error:
        if "dedicated app/resource prefix" in str(error):
            raise NativeToolUnavailable(str(error)) from error
        raise
    before_binding = retain_native_reader_runtime(root, runtime_before, run_id=run_id)
    parse_native_bytes(kind, artifact.content)
    source = directory / f"input.{kind}"
    destination = directory / "input.pdf"
    profile = directory / "office-profile"
    if any(path.exists() for path in (source, destination, profile)):
        raise ValueError("native Office conversion requires fresh input/output/profile")
    atomic_write_bytes_confined(root, source, artifact.content, mode=0o600)
    reader = SnapshotReader(root)
    consumed = reader.read(source.relative_to(root).as_posix())
    if consumed.sha256 != artifact.sha256:
        raise ValueError("Office input differs from admitted bytes")
    version = native_tool_record(
        root, directory, "soffice", ["--version"], env=env, deadline=deadline, run_id=run_id, purpose="office-version"
    )
    filter_name = "writer_pdf_Export" if kind == "docx" else "impress_pdf_Export"
    tool = native_tool_record(
        root,
        directory,
        "soffice",
        [
            f"-env:UserInstallation={profile.as_uri()}",
            "--headless",
            "--nologo",
            "--nodefault",
            "--nolockcheck",
            "--nofirststartwizard",
            "--convert-to",
            f"pdf:{filter_name}",
            "--outdir",
            str(directory),
            str(source),
        ],
        env=env,
        deadline=deadline,
        run_id=run_id,
        purpose="office-conversion",
    )
    if any(record["exit_code"] != 0 or record["original_exit_code"] != 0 for record in (version, tool)):
        raise ValueError("actual native Office conversion/version probe failed")
    runtime_after = observe_native_reader_runtime(Path(selected), deadline=deadline)
    version_raw = reader.read(version["artifact"]["path"]).json_object()
    validate_native_reader_runtime(runtime_before, runtime_after, executable_sha256=version_raw["executable_sha256"])
    after_binding = retain_native_reader_runtime(root, runtime_after, run_id=run_id)
    converted = reader.read(destination.relative_to(root).as_posix())
    observation = parse_native_bytes("pdf", converted.content)
    if not reader.unchanged():
        raise ValueError("native Office input/output changed during inspection")
    return {
        "schema_version": "template-native-office-conversion-v2",
        "kind": kind,
        "input": {"path": consumed.relative, "sha256": consumed.sha256},
        "pdf": {"path": converted.relative, "sha256": converted.sha256},
        "observation": observation,
        "tool_artifacts": [version["artifact"], tool["artifact"]],
        "reader_runtime_before": before_binding,
        "reader_runtime_after": after_binding,
        "status": "blocked",
        "blockers": ["complete source/layout and valid visual controls required"],
    }


def observe_office_layout(
    root: Path,
    directory: Path,
    converted: dict[str, Any],
    *,
    source_contract: NativeContentContract | None = None,
    projection_pages: tuple[ProjectionPage, ...] = (),
    env: dict[str, str],
    deadline: float,
    run_id: str,
) -> dict[str, Any]:
    """Inspect all native-converted pages against independent source obligations."""
    from infrastructure.reporting.native_format_acceptance_checks import audit_pdf_raster
    from infrastructure.reporting.native_format_pdf_oracles import inspect_pdf_content
    from infrastructure.reporting.native_projection_checks import inspect_projection

    reader = SnapshotReader(root)
    artifact = reader.read(converted["pdf"]["path"])
    if artifact.sha256 != converted["pdf"]["sha256"]:
        raise ValueError("native Office PDF differs from admitted conversion receipt")
    parsed = parse_native_bytes("pdf", artifact.content)
    raster = audit_pdf_raster(root, directory, artifact, parsed, env=env, deadline=deadline, run_id=run_id, ordinal=0)
    pixels = tuple(reader.read(row["path"]).content for row in raster["rasters"])
    if converted["kind"] == "pptx":
        if not projection_pages:
            raise ValueError("PPTX layout requires complete independently reconstructed source slides")
        observation = inspect_projection(
            artifact.content, pixels, projection_pages, deadline=deadline, native_context=(root, directory, env, run_id)
        )
    else:
        if source_contract is None:
            raise ValueError("DOCX layout requires independently reconstructed source contract")
        observation = inspect_pdf_content(artifact.content, pixels, source_contract, deadline=deadline)
    if not reader.unchanged():
        raise ValueError("actual Office converted PDF/raster custody changed")
    return {"conversion": converted, "raster": raster, "observation": observation, "status": observation["status"]}


OFFICE_CONTROL_CODES = {
    "baseline": None,
    "changed-number": "pdf.content.missing",
    "tiny-text": "pdf.layout.font",
    "white-text": "pdf.visual.invisible",
    "reordered-text": "pdf.content.order",
    "clipped-text": "pdf.layout.crop",
    "opaque-overlay": "pdf.visual.invisible",
}
PPTX_CONTROL_CODES = {
    "baseline": None,
    "changed-number": "pdf.content.missing",
    "tiny-text": "projection.typography.font",
    "clipped-text": "pdf.layout.crop",
    "opaque-overlay": "pdf.visual.invisible",
    "wrong-citation": "projection.citation.target",
}


def reconstruct_office_control_fixture(
    root: Path, directory: Path, kind: str, name: str, *, env: dict[str, str], deadline: float, run_id: str
) -> tuple[bytes, NativeContentContract | None, tuple[ProjectionPage, ...], dict[str, Any] | None]:
    """Reconstruct exact source-owned fixture with real Pandoc where required."""
    from io import BytesIO
    import zipfile
    from defusedxml.ElementTree import fromstring
    from xml.etree.ElementTree import SubElement, tostring
    from infrastructure.reporting.native_format_content_contracts import content_contract_from_bytes
    from infrastructure.reporting.native_projection_contracts import compile_projection_contract
    from infrastructure.rendering.slide_deck import DeckContent, Slide
    from infrastructure.rendering.pptx_deck import render_pptx

    codes = OFFICE_CONTROL_CODES if kind == "docx" else PPTX_CONTROL_CODES if kind == "pptx" else {}
    if name not in codes:
        raise ValueError("unknown native Office control")
    source_contract: NativeContentContract | None = None
    projection_pages: tuple[ProjectionPage, ...] = ()
    tool: dict[str, Any] | None = None
    original = directory / f"fixture.{kind}"
    if kind == "docx":
        source = b"# Invented measurement\n\nMeasured value -3.25 mV.\n\nIndependent repetition -3.25 mV.\n"
        source_path = directory / "fixture.md"
        producer_source = source
        if name == "opaque-overlay":
            from PIL import Image

            pixels = BytesIO()
            Image.new("RGB", (64, 64), "white").save(pixels, format="PNG")
            image_path = directory / "native-overlay.png"
            atomic_write_bytes_confined(root, image_path, pixels.getvalue(), mode=0o600)
            producer_source += b"\n![](" + str(image_path).encode() + b"){width=8.5in height=4.16667in}\n"
        atomic_write_bytes_confined(root, source_path, producer_source, mode=0o600)
        source_contract = content_contract_from_bytes("office:docx-control", (("invented.md", source),))
        tool = native_tool_record(
            root,
            directory,
            "pandoc",
            [str(source_path), "--standalone", "-o", str(original)],
            env=env,
            deadline=deadline,
            run_id=run_id,
            purpose="docx-control-producer",
        )
        if tool["exit_code"] != 0 or tool["original_exit_code"] != 0:
            raise ValueError("actual DOCX control producer failed")
        namespace = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
        with zipfile.ZipFile(original) as archive:
            members = [(info, archive.read(info.filename)) for info in archive.infolist()]
        output = BytesIO()
        for info, content in members:
            if info.filename == "word/document.xml" and name != "baseline":
                document = fromstring(content)
                if name == "changed-number":
                    for text in document.iter(namespace + "t"):
                        if text.text and "-3.25" in text.text:
                            text.text = text.text.replace("-3.25", "-3.26", 1)
                            break
                elif name in {"tiny-text", "white-text"}:
                    for run in document.iter(namespace + "r"):
                        if not any(text.text and "-3.25" in text.text for text in run.iter(namespace + "t")):
                            continue
                        properties = run.find(namespace + "rPr")
                        if properties is None:
                            properties = SubElement(run, namespace + "rPr")
                            run.remove(properties)
                            run.insert(0, properties)
                        tag = "sz" if name == "tiny-text" else "color"
                        existing = properties.find(namespace + tag)
                        if existing is None:
                            existing = SubElement(properties, namespace + tag)
                        existing.set(namespace + "val", "8" if name == "tiny-text" else "FFFFFF")
                elif name == "clipped-text":
                    body = document.find(namespace + "body")
                    if body is None:
                        raise ValueError("DOCX clipped control has no body")
                    paragraph = list(body.findall(namespace + "p"))[1]
                    properties = paragraph.find(namespace + "pPr")
                    if properties is None:
                        properties = SubElement(paragraph, namespace + "pPr")
                        paragraph.remove(properties)
                        paragraph.insert(0, properties)
                    frame = SubElement(properties, namespace + "framePr")
                    for key, value in {
                        "hAnchor": "page",
                        "vAnchor": "page",
                        "x": "-1440",
                        "y": "1800",
                        "w": "10000",
                        "wrap": "none",
                    }.items():
                        frame.set(namespace + key, value)
                elif name == "opaque-overlay":
                    wp = "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}"
                    replaced = 0
                    for parent in document.iter():
                        for inline in list(parent):
                            if inline.tag != wp + "inline":
                                continue
                            anchor = fromstring(OFFICE_OVERLAY_ANCHOR_XML)
                            for child in list(inline):
                                anchor.append(child)
                                if child.tag == wp + "extent":
                                    SubElement(anchor, wp + "wrapNone")
                            parent.remove(inline)
                            parent.append(anchor)
                            replaced += 1
                    if replaced != 1:
                        raise ValueError("actual Pandoc overlay fixture must contain one known-valid image")
                elif name == "reordered-text":
                    body = document.find(namespace + "body")
                    if body is None:
                        raise ValueError("DOCX control fixture has no body")
                    paragraphs = list(body.findall(namespace + "p"))
                    if len(paragraphs) != 3:
                        raise ValueError("DOCX control fixture paragraphs differ from reviewed baseline")
                    first, second = paragraphs[1:]
                    first_index = list(body).index(first)
                    body.remove(first)
                    body.remove(second)
                    body.insert(first_index, second)
                    body.insert(first_index + 1, first)
                content = tostring(document, encoding="utf-8", xml_declaration=True)
            with zipfile.ZipFile(output, "a") as target:
                target.writestr(info, content)
        atomic_write_bytes_confined(root, original, output.getvalue(), mode=0o600)
    else:
        from pptx import Presentation
        from pptx.util import Pt
        from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
        from pptx.dml.color import RGBColor

        deck = DeckContent(
            "Invented measurement",
            slides=(
                Slide("Invented measurement", kind="title"),
                Slide("Actual measurement", bullets=("Measured value -3.25 mV.",), source="invented.csv"),
            ),
        )
        projection_pages = compile_projection_contract(deck, source_base_url="https://example.org")
        render_pptx(deck, original, source_base_url="https://example.org")
        presentation = Presentation(str(original))
        for shape in presentation.slides[1].shapes:
            if not shape.has_text_frame:
                continue
            if "Measured value" in shape.text:
                if name == "changed-number":
                    for paragraph in shape.text_frame.paragraphs:
                        for run in paragraph.runs:
                            run.text = run.text.replace("-3.25", "-3.26")
                elif name == "tiny-text":
                    for paragraph in shape.text_frame.paragraphs:
                        paragraph.font.size = Pt(5)
                        for run in paragraph.runs:
                            run.font.size = Pt(5)
                elif name == "clipped-text":
                    shape.left = Pt(-60)
                elif name == "opaque-overlay":
                    overlay = presentation.slides[1].shapes.add_shape(
                        MSO_AUTO_SHAPE_TYPE.RECTANGLE,
                        shape.left - Pt(9),
                        shape.top - Pt(18),
                        shape.width + Pt(18),
                        shape.height + Pt(36),
                    )
                    overlay.fill.solid()
                    overlay.fill.fore_color.rgb = RGBColor(255, 255, 255)
                    overlay.line.fill.background()
            if "invented.csv" in shape.text and name == "wrong-citation":
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        run.hyperlink.address = "https://example.org/wrong"
        presentation.save(str(original))
    fixture = SnapshotReader(root).read(original.relative_to(root).as_posix())
    return fixture.content, source_contract, projection_pages, tool


def run_office_control(
    root: Path, directory: Path, kind: str, name: str, *, env: dict[str, str], deadline: float, run_id: str
) -> dict[str, Any]:
    """Generate actual valid Office packages and test real native conversion mutants."""
    _, source_contract, projection_pages, _ = reconstruct_office_control_fixture(
        root, directory, kind, name, env=env, deadline=deadline, run_id=run_id
    )
    original = directory / f"fixture.{kind}"
    codes = OFFICE_CONTROL_CODES if kind == "docx" else PPTX_CONTROL_CODES
    reader = SnapshotReader(root)
    artifact = reader.read(original.relative_to(root).as_posix())
    parsed = parse_native_bytes(kind, artifact.content)
    conversion = convert_office(root, directory, artifact, kind, env=env, deadline=deadline, run_id=run_id)
    result = observe_office_layout(
        root,
        directory,
        conversion,
        source_contract=source_contract,
        projection_pages=projection_pages,
        env=env,
        deadline=deadline,
        run_id=run_id,
    )
    expected = codes[name]
    observation = result["observation"]
    findings = {finding["code"] for finding in observation["findings"]}
    if (
        observation["blockers"]
        or (expected is None and observation["status"] != "pass")
        or (expected is not None and (observation["status"] != "fail" or expected not in findings))
    ):
        raise ValueError("actual valid Office control did not produce intended source/layout result")
    return {
        "schema_version": "template-native-office-valid-control-v1",
        "kind": kind,
        "name": name,
        "expected_code": expected,
        "parser": parsed,
        "layout": result,
        "oracle_exit_code": 0 if expected is None else 1,
    }


OFFICE_OVERLAY_ANCHOR_XML = b"""<wp:anchor
 xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
 distT="0" distB="0" distL="0" distR="0" simplePos="0"
 relativeHeight="251658240" behindDoc="0" locked="0" layoutInCell="1" allowOverlap="1">
 <wp:simplePos x="0" y="0"/><wp:positionH relativeFrom="page"><wp:posOffset>0</wp:posOffset></wp:positionH>
 <wp:positionV relativeFrom="page"><wp:posOffset>1270000</wp:posOffset></wp:positionV>
 </wp:anchor>"""


def supervise_office_controls(
    root: Path, directory: Path, kind: str, *, interpreter: Path, env: dict[str, str], deadline: float, run_id: str
) -> list[dict[str, Any]]:
    """Require actual baseline-first native conversions and original mutant exits."""
    import time
    import json
    from infrastructure.core.execution_boundary import run_bounded_subprocess
    from infrastructure.project.public_acceptance_evidence import bootstrap_command, write_json
    from infrastructure.reporting.native_format_acceptance_contracts import create_private_native_directory
    from infrastructure.reporting.native_format_acceptance_checks import NativeToolUnavailable

    codes = OFFICE_CONTROL_CODES if kind == "docx" else PPTX_CONTROL_CODES if kind == "pptx" else {}
    if not codes:
        raise ValueError("supervised Office controls require DOCX/PPTX")
    records = []
    for ordinal, (name, diagnostic) in enumerate(codes.items()):
        child = directory / f"{kind}-control-{ordinal}"
        create_private_native_directory(root, child)
        command = bootstrap_command(
            interpreter,
            root,
            __name__,
            "--kind",
            kind,
            "--control",
            name,
            "--root",
            str(root),
            "--directory",
            str(child),
            "--deadline",
            str(deadline),
            "--run-id",
            run_id,
        )
        remaining = deadline - time.monotonic()
        if remaining <= 0:
            raise TimeoutError("Office controls exhausted original deadline")
        process = run_bounded_subprocess(command, cwd=root, env=env, timeout=min(remaining, 120))
        receipt = write_json(
            root,
            directory / f"{kind}-control-{ordinal}-process.json",
            {
                "schema_version": "template-native-office-control-process-v1",
                "run_id": run_id,
                "command": list(command),
                "exit_code": process.returncode,
                "original_exit_code": process.original_returncode,
                "timed_out": process.timed_out,
                "custody_refused": process.custody_refused,
                "stdout": process.stdout,
                "stderr": process.stderr,
                "diagnostic": process.command_error,
            },
        )
        if process.returncode == 75:
            raise NativeToolUnavailable(f"actual {kind} control {name} requires missing native engine")
        expected = 0 if diagnostic is None else 1
        if (
            process.timed_out
            or process.custody_refused
            or process.returncode != expected
            or process.original_returncode != expected
        ):
            raise ValueError("actual native Office control did not preserve intended original outcome")
        result = SnapshotReader(root).read((child / "observation.json").relative_to(root).as_posix())
        payload = json.loads(result.content)
        observed = payload["layout"]["observation"]
        if (
            payload["kind"] != kind
            or payload["name"] != name
            or payload["expected_code"] != diagnostic
            or observed["blockers"]
        ):
            raise ValueError("actual Office control receipt differs from source-owned obligations")
        if diagnostic is not None and diagnostic not in {item["code"] for item in observed["findings"]}:
            raise ValueError("actual Office mutant has no intended native diagnostic")
        records.append(
            {
                "kind": kind,
                "name": name,
                "process_artifact": receipt,
                "artifact": {"path": result.relative, "sha256": result.sha256},
                "original_exit_code": expected,
            }
        )
    return records


def main() -> int:
    """Run one actual native Office control under its supervisor's original deadline."""
    import argparse
    import os
    from infrastructure.reporting.native_format_acceptance_checks import NativeToolUnavailable
    from infrastructure.project.public_acceptance_evidence import write_json

    parser = argparse.ArgumentParser()
    parser.add_argument("--kind", choices=("docx", "pptx"), required=True)
    parser.add_argument("--control", required=True)
    parser.add_argument("--root", type=Path, required=True)
    parser.add_argument("--directory", type=Path, required=True)
    parser.add_argument("--deadline", type=float, required=True)
    parser.add_argument("--run-id", required=True)
    args = parser.parse_args()
    try:
        result = run_office_control(
            args.root,
            args.directory,
            args.kind,
            args.control,
            env=dict(os.environ),
            deadline=args.deadline,
            run_id=args.run_id,
        )
    except NativeToolUnavailable:
        return 75
    write_json(args.root, args.directory / "observation.json", result)
    return int(result["oracle_exit_code"])


if __name__ == "__main__":
    raise SystemExit(main())

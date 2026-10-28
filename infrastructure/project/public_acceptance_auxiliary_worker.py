"""Real source-owned auxiliary checks executed only in a supervised child."""

from __future__ import annotations

import argparse
import importlib
import json
import os
import platform
import subprocess
import sys
import tempfile
import time
import tomllib
from pathlib import Path
from typing import Any
from collections.abc import Callable

from infrastructure.project.public_acceptance_evidence import MissingAuxiliaryProvider, read_json, write_json
from infrastructure.reporting.native_format_acceptance_checks import NativeToolUnavailable
from infrastructure.reporting.native_format_acceptance_contracts import NativeDeadlineExpired, NativeDeadlineUnavailable


def execute(
    root: Path,
    gate: str,
    budget: int,
    *,
    directory: Path | None = None,
    inputs: dict[str, Any] | None = None,
    operation_deadline: float | None = None,
) -> dict[str, Any]:
    """Invoke actual owned implementations; never substitute fixture success."""
    if gate == "root:roster":
        import hashlib
        from infrastructure.project.public_scope import PUBLIC_PROJECT_NAMES, public_project_source_roots
        from infrastructure.reporting.release_readiness_snapshot import SnapshotReader

        roots = public_project_source_roots(root)
        observed = [path.parent.relative_to(root / "projects").as_posix() for path in roots]
        actual_directories = sorted(
            "templates/" + path.name for path in (root / "projects/templates").iterdir() if path.is_dir()
        )
        if observed != list(PUBLIC_PROJECT_NAMES) or actual_directories != sorted(observed):
            raise ValueError("actual public source roster differs from declared complete roster")
        reader = SnapshotReader(root)
        observations = []
        for scope in observed:
            snapshot = reader.read(f"projects/{scope}/pyproject.toml")
            metadata = tomllib.loads(snapshot.content.decode())["project"]
            observations.append(
                {
                    "scope": scope,
                    "package": metadata["name"],
                    "version": metadata["version"],
                    "metadata_artifact": {
                        "path": snapshot.relative,
                        "sha256": hashlib.sha256(snapshot.content).hexdigest(),
                    },
                }
            )
        if not reader.unchanged():
            raise ValueError("actual roster metadata changed during observation")
        return {"projects": observed, "source_roster_observations": observations}
    if gate == "root:docs_lint":
        from infrastructure.validation.docs import lint_runner as lint

        result: dict[str, Any] = {}
        checks: dict[str, Callable[..., Any]] = {
            "mermaid": lint.run_mermaid_lint,
            "broken_links": lint.run_links_lint,
            "consistency": lint.run_consistency_lint,
            "doc_pairs": lint.run_doc_pairs_lint,
        }
        for name, check in checks.items():
            try:
                result[name] = [str(issue) for issue in check(root, quiet=True)]
            except RuntimeError as error:
                raise MissingAuxiliaryProvider(str(error)) from error
        if any(result.values()):
            raise ValueError(json.dumps(result))
        return result
    if gate == "root:methods":
        from infrastructure.methods.orchestration import audit_methods_projects
        from infrastructure.project.public_scope import PUBLIC_PROJECT_NAMES

        report = audit_methods_projects(root, list(PUBLIC_PROJECT_NAMES), artifact_mode="rendered")
        if not report.passed or len(report.projects) != len(PUBLIC_PROJECT_NAMES):
            raise ValueError(json.dumps(report.to_dict()))
        return {"assertions": {"complete_rendered_methods_audit": True}, "methods_report": report.to_dict()}
    if gate == "root:macos:process":
        if platform.system() != "Darwin":
            raise MissingAuxiliaryProvider("macOS process acceptance requires an actual Darwin host")
        # Verify ordinary exit transport and a real OS timeout, rather than
        # deriving platform acceptance from uname or a structural marker.
        child = subprocess.run([sys.executable, "-I", "-B", "-c", "raise SystemExit(19)"], check=False, timeout=5)
        if child.returncode != 19:
            raise ValueError("actual process exit transport changed")
        try:
            subprocess.run([sys.executable, "-I", "-B", "-c", "import time; time.sleep(30)"], check=False, timeout=0.2)
        except subprocess.TimeoutExpired:
            result = {"platform": "macos", "assertions": {"exit_transport": True, "real_timeout": True}}
            if directory is not None:
                import psutil
                from infrastructure.project.public_acceptance_auxiliary import supervise

                command = [
                    sys.executable,
                    "-I",
                    "-B",
                    "-c",
                    "import subprocess,sys,time; "
                    "p=subprocess.Popen([sys.executable,'-I','-B','-c','import time; time.sleep(30)']); "
                    "print(p.pid,flush=True); time.sleep(30)",
                ]
                # This intentional timeout must first admit Source and ARM the
                # independent guardian. It remains inside the enclosing gate's
                # original clock; a startup refusal proves no descendant cleanup.
                process = supervise(command, root, directory / "descendant", time.monotonic() + min(10, budget))
                if type(process.get("pid")) is not int or type(process.get("exit_code")) is not int:
                    raise ValueError("macOS descendant control did not launch and observe an actual child")
                pid = int((root / process["streams"]["stdout"]["path"]).read_text().strip())
                try:
                    descendant = psutil.Process(pid)
                    alive = descendant.is_running() and descendant.status() != psutil.STATUS_ZOMBIE
                except psutil.NoSuchProcess:
                    alive = False
                if not process["timed_out"] or alive:
                    raise ValueError("actual timed-out process group left a running descendant")
                result["assertions"]["descendant_cleanup"] = True
                result["descendant_process"] = process
            return result
        raise ValueError("actual sleeping child did not time out")
    if gate == "root:macos:rendering":
        if platform.system() != "Darwin":
            raise MissingAuxiliaryProvider("macOS rendering acceptance requires an actual Darwin host")
        import shutil

        for name in ("pandoc", "pdftoppm", "pdffonts"):
            if shutil.which(name) is None:
                raise MissingAuxiliaryProvider(f"{name} is unavailable on the actual Darwin host")
        with tempfile.TemporaryDirectory(prefix="template-macos-render-") as temporary:
            output = Path(temporary) / "render.html"
            rendered = subprocess.run(
                ["pandoc", "--from", "markdown", "--to", "html5", "--output", str(output)],
                input="# Acceptance\n\nA real **render**.\n",
                text=True,
                capture_output=True,
                timeout=min(budget, 30),
                check=False,
            )
            content = output.read_text() if output.is_file() else ""
            if (
                rendered.returncode != 0
                or "<strong>render</strong>" not in content
                or '<h1 id="acceptance">' not in content
            ):
                raise ValueError("actual Darwin pandoc rendering failed its content oracle")
            from infrastructure.reporting.native_print_witness import run_print_witness

            if directory is None:
                raise ValueError("original Darwin print witness requires a confined evidence directory")
            witness = run_print_witness(root, directory / "print-witness", budget)
            return {
                "platform": "macos",
                "assertions": {"actual_pandoc_content": True},
                "rendered_html": content,
                "print_witness_artifact": witness,
            }
    if gate.endswith(":pipeline"):
        from infrastructure.project.public_acceptance_auxiliary_pipeline import run_pipeline

        if directory is None:
            raise ValueError("actual portable pipeline requires a confined evidence directory")
        return run_pipeline(
            root, gate.removesuffix(":pipeline"), directory, budget, operation_deadline=operation_deadline
        )
    if directory is not None:
        from infrastructure.project.public_acceptance_auxiliary_native import native_execution

        native = native_execution(root, gate, budget, directory, inputs or {}, operation_deadline=operation_deadline)
        if native is not None:
            return native
    scope, kind = gate.rsplit(":", 1)
    if ":live:" in gate:
        scope, capability = gate.split(":live:", 1)
        kind = f"live:{capability}"
    project = root / "projects" / scope
    metadata = tomllib.loads((project / "pyproject.toml").read_text())
    provider = metadata.get("tool", {}).get("template", {}).get("acceptance_producers", {}).get(kind)
    if not isinstance(provider, str) or provider.count(":") != 1:
        raise MissingAuxiliaryProvider(f"{scope} has no source-owned acceptance_producers declaration for {kind}")
    module_name, callable_name = provider.split(":")
    if not all(part.isidentifier() for part in module_name.split(".")) or not callable_name.isidentifier():
        raise ValueError("invalid source-owned producer reference")
    sys.path.insert(0, str(project / "src"))
    module = importlib.import_module(module_name)
    source = Path(module.__file__ or "").absolute()
    if not source.resolve(strict=True).is_relative_to((project / "src").resolve()) or source.is_symlink():
        raise ValueError("producer did not originate in the selected public project")
    value = getattr(module, callable_name)(project_root=project, timeout_seconds=budget)
    if not isinstance(value, dict):
        raise ValueError("actual producer did not return its raw observations")
    return value


def negative_control(root: Path, gate: str) -> int:
    """Run the same production checker on an independently invalid input."""
    if gate.endswith(":pipeline"):
        from infrastructure.project.public_acceptance_auxiliary_pipeline import SCHEMA, validate_portable_pipeline
        from infrastructure.reporting.release_readiness_snapshot import SnapshotReader

        try:
            validate_portable_pipeline(
                gate.removesuffix(":pipeline"), {"pipeline_schema": SCHEMA}, SnapshotReader(root)
            )
        except ValueError:
            return 23
        return 0
    if gate == "root:docs_lint":
        from infrastructure.validation.docs.cross_link_lint import find_broken_links

        with tempfile.TemporaryDirectory(prefix="template-doc-negative-") as temporary:
            path = Path(temporary) / "README.md"
            path.write_text("[broken](missing-file.md)\n")
            return 23 if find_broken_links([path]) else 0
    if gate == "root:methods":
        from infrastructure.methods.orchestration import audit_methods_projects

        with tempfile.TemporaryDirectory(prefix="template-method-negative-") as temporary:
            report = audit_methods_projects(Path(temporary), ["templates/missing"], artifact_mode="rendered")
            return 23 if not report.passed else 0
    if gate == "root:macos:rendering":
        child = subprocess.run(
            ["pandoc", "--from", "invalid-template-format"], capture_output=True, timeout=10, check=False
        )
        return 23 if child.returncode != 0 else 0
    if gate == "root:macos:process":
        child = subprocess.run([sys.executable, "-I", "-B", "-c", "raise SystemExit(23)"], timeout=5, check=False)
        return child.returncode
    scope = gate.split(":", 1)[0]
    from infrastructure.project.public_acceptance_auxiliary_native import project_module

    if gate.endswith(":domain"):
        from infrastructure.project.public_acceptance_auxiliary_domains import domain_negative

        observed = domain_negative(root, scope)
        if observed is not None:
            return observed
    if scope == "templates/template_data_descriptor":
        checker = project_module(root, scope, "data_descriptor.verification")
        descriptor = json.loads((root / "projects" / scope / "data/example_descriptor.json").read_text())
        descriptor["files"][0]["checksum"] = "sha256:" + "0" * 64
        observations = checker.verify_descriptor_files(descriptor, root / "projects" / scope / "data")
        return 23 if observations[0].status == "checksum_mismatch" else 0
    if scope == "templates/template_methods_paper":
        units = project_module(root, scope, "template_methods_paper.methods_dsl.units")
        try:
            units.Quantity(1, "mL") + units.Quantity(1, "g")
        except units.DimensionError:
            return 23
        return 0
    if scope == "templates/template_sia":
        domain = project_module(root, scope, "template_sia.loop.proposal_domain")
        try:
            domain.validated_threshold(float("nan"))
        except ValueError:
            return 23
        return 0
    if scope == "templates/template_autoscientists":
        domain = project_module(root, scope, "template_autoscientists.agents.live_domain")
        try:
            domain.parse_proposal('{"axis":0,"step":2}', (0,), "control", frozenset())
        except (ValueError, RuntimeError):
            return 23
        return 0
    if scope == "templates/template_autopoiesis":
        checker = project_module(root, scope, "template_autopoiesis.gates.child_results")
        with tempfile.TemporaryDirectory(prefix="template-child-negative-") as temporary:
            child_root = Path(temporary)
            path = child_root / "output/data/analysis_results.json"
            path.parent.mkdir(parents=True)
            path.write_text(
                json.dumps({"schema_version": "autopoiesis/child-analysis/1", "domain": "optimization", "results": {}})
            )
            try:
                checker.verify_analysis_results(child_root, "optimization")
            except ValueError:
                return 23
        return 0
    if scope == "templates/template_active_inference":
        if gate.endswith(":lean"):
            checker = project_module(root, scope, "gates.lean")
            passed, _ = checker.lean_axioms_clean(root / "projects" / scope, declarations=("invalid;declaration",))
            return 0 if passed else 23
        import numpy as np

        checker = project_module(root, scope, "simulation.si_loop")
        try:
            checker._normalise_sampling_distribution(np.array([-1.0, 2.0]), label="negative-control")
        except ValueError:
            return 23
        return 0
    if scope == "templates/template_formal" and (gate.endswith(":network-faults") or gate.endswith(":domain")):
        checker = project_module(root, scope, "template_formal.network.bus")
        protocol = project_module(root, scope, "template_formal.protocol.session")
        ids = project_module(root, scope, "template_formal.types.ids")
        bus = checker.InProcessBus(encode=protocol.encode_wire_message, decode=protocol.decode_wire_message)
        try:
            bus.try_recv(ids.new_agent_id())
        except checker.UnknownEndpointError:
            return 23
        return 0
    if scope == "templates/template_pools_rules_tools":
        checker = project_module(root, scope, "template_pools_rules_tools.resource_paths")
        try:
            checker.public_resource_root("private")
        except ValueError:
            return 23
        return 0
    if ":live:local-review" in gate or ":live:translation" in gate:
        from infrastructure.llm.core.local_ollama import validate_local_endpoint

        try:
            validate_local_endpoint("https://unapproved.invalid")
        except ValueError:
            return 23
        return 0
    if (
        any(gate.endswith(suffix) for suffix in (":print-raster", ":raster-redaction", ":pdf-pptx"))
        or gate.endswith(":domain")
        and scope
        in {
            "templates/template_newspaper",
            "templates/template_storybook",
            "templates/template_redacted_report",
            "templates/template_pitch_deck",
        }
    ):
        from pypdf import PdfReader
        from pypdf.errors import PdfReadError
        import io

        try:
            PdfReader(io.BytesIO(b"%PDF-1.7\ninvalid-publication\n"))
        except PdfReadError:
            return 23
        return 0
    if scope == "templates/template_eda_notebook":
        import pandas as pd
        from infrastructure.project.public_acceptance_auxiliary_native import project_module

        oracle = project_module(root, scope, "template_eda_notebook.eda.notebook_oracle")
        raw = pd.read_csv(root / "projects" / scope / "data/measurements.csv")
        incorrect = oracle.reference_observations(raw)
        incorrect["summaries"][0]["mean"] += 100
        return 23 if oracle.verify_observations(incorrect, raw) else 0
    from infrastructure.project.public_acceptance_auxiliary_native import LITERATURE_MODULES

    if scope in LITERATURE_MODULES:
        from infrastructure.search.literature.live_contract import LiteratureFailed
        from infrastructure.search.literature.live_providers import parse_licensed_jats

        try:
            parse_licensed_jats(b"<article><body>unlicensed wrong-identity content</body></article>", "10.1/wrong")
        except LiteratureFailed:
            return 23
        return 0
    raise MissingAuxiliaryProvider("domain/live producer must supply its own executed negative-control artifacts")


def main() -> int:
    """Write raw facts before exiting, including the true failure class."""
    parser = argparse.ArgumentParser()
    parser.add_argument("--spec", type=Path, required=True)
    args = parser.parse_args()
    root = Path.cwd().resolve()
    if not Path(__file__).resolve().is_relative_to(root / "infrastructure/project"):
        raise ValueError("auxiliary worker implementation was not loaded from the nominated source root")
    return run_specification(root, read_json(root, args.spec))


def run_specification(root: Path, spec: dict[str, Any]) -> int:
    """Retain true terminal facts from the actual selected source-owned check."""
    result: dict[str, Any] = {"gate": spec["gate"], "pid": os.getpid(), "platform": platform.system()}
    result["runtime"] = {
        "executable": sys.executable,
        "python_version": platform.python_version(),
        "implementation": platform.python_implementation(),
        "pycache_prefix": sys.pycache_prefix,
    }
    from infrastructure.project.public_acceptance_auxiliary_pipeline import AuxiliaryPipelineFailure

    try:
        if spec.get("negative"):
            code = negative_control(root, spec["gate"])
            result.update(exit_code=code, status="negative-control")
        else:
            result.update(
                details=execute(
                    root,
                    spec["gate"],
                    spec["budget"],
                    directory=(root / spec["raw_path"]).parent,
                    inputs=spec.get("producer_inputs", {}),
                    operation_deadline=spec.get("operation_deadline"),
                ),
                status="pass",
            )
            code = 0
    except AuxiliaryPipelineFailure as error:
        result.update(error.worker_observation)
        code = error.code
    except MissingAuxiliaryProvider as error:
        result.update(status="blocked", diagnostic=str(error))
        code = 75
    except NativeDeadlineExpired as error:
        result.update(status="timeout", diagnostic=str(error), failure_class=type(error).__name__)
        code = 124
    except (
        OSError,
        ValueError,
        TypeError,
        KeyError,
        subprocess.SubprocessError,
        NativeToolUnavailable,
        NativeDeadlineUnavailable,
    ) as error:
        result.update(status="fail", diagnostic=str(error), failure_class=type(error).__name__)
        code = 1
    result["exit_code"] = code
    write_json(root, root / spec["raw_path"], result)
    return code


def publication_native(root: Path, scope: str, capability: str, directory: Path) -> dict[str, Any]:
    """Render actual source content in a disposable copy and inspect output pixels."""
    from infrastructure.project.public_acceptance_auxiliary_native import project_module, native_result
    import io
    import logging
    from pypdf import PdfReader
    import pypdfium2
    from PIL import ImageStat
    from infrastructure.project.installed_acceptance_artifacts import write_binary

    project = root / "projects" / scope
    with tempfile.TemporaryDirectory(prefix="template-auxiliary-publication-") as temporary:
        copied = Path(temporary) / project.name
        from infrastructure.reporting.readiness_receipts import public_source_paths
        from infrastructure.reporting.release_readiness_snapshot import SnapshotReader

        copied.mkdir()
        source_reader = SnapshotReader(root)
        for path in public_source_paths(root, scope):
            relative = path.relative_to(project)
            target = copied / relative
            target.parent.mkdir(parents=True, exist_ok=True)
            target.write_bytes(source_reader.read(path.relative_to(root).as_posix()).content)
        if not source_reader.unchanged():
            raise ValueError("actual publication source changed during confined public export")
        pptx_bindings = {}
        if scope == "templates/template_newspaper":
            module = project_module(root, scope, "newspaper.engine")
            result = module.build_and_render(copied)
            if not result.all_pages_fit:
                raise ValueError("actual newspaper print layout overflowed")
            pdfs = [result.output_path]
        elif scope == "templates/template_storybook":
            module = project_module(root, scope, "storybook.rendering")
            result = module.build_storybook_pdf(copied)
            pdfs = [result.output_path]
        elif scope == "templates/template_pitch_deck":
            module = project_module(root, scope, "template_pitch_deck.render_orchestration")
            paths = module.render_all_decks(copied, root, logging.getLogger(__name__))
            pdfs = [path for path in paths if path.suffix == ".pdf"]
            from pptx import Presentation

            pptxs = [path for path in paths if path.suffix == ".pptx"]
            if not pdfs or len(pptxs) != len(pdfs):
                raise ValueError("actual deck producer omitted a declared PDF/PPTX pair")
            for pdf in pdfs:
                match = next((path for path in pptxs if path.stem == pdf.stem), None)
                if match is None or len(Presentation(str(match)).slides) != len(PdfReader(pdf).pages):
                    raise ValueError("actual PDF/PPTX slide inventories differ")
                pptx_bindings[pdf.name] = write_binary(root, directory / match.name, match.read_bytes())
        else:
            module = project_module(root, scope, "redacted_report.visuals")
            redaction = project_module(root, scope, "redacted_report.redaction")
            fixture = json.loads((copied / "data/example_segments.json").read_text())
            segments = [redaction.RedactionSegment(**item) for item in fixture["segments"]]
            decisions = [redaction.RedactionDecision(**item) for item in fixture["redactions"]]
            output = copied / "output/auxiliary-redaction"
            module.write_dev_variant_pdfs(segments, decisions, output, include_steganography=False, include_kmyth=False)
            check = module.verify_dev_variant_outputs(output, render_smoke=True, require_steganography=False)
            if check.get("valid") is not True:
                raise ValueError("actual redaction proof matrix failed: " + str(check))
            pdfs = sorted(output.glob("*.pdf"))
        artifacts = []
        for index, pdf in enumerate(pdfs):
            pages = len(PdfReader(pdf).pages)
            if not 0 < pages <= 200:
                raise ValueError("actual print PDF page inventory is empty or exceeds the ceiling")
            with pypdfium2.PdfDocument(pdf) as document:
                raster = []
                for number in range(pages):
                    image = document[number].render(scale=0.5).to_pil()
                    if image.width * image.height > 20_000_000 or max(ImageStat.Stat(image.convert("L")).stddev) <= 0:
                        raise ValueError("actual print raster is oversized or blank")
                    buffer = io.BytesIO()
                    image.save(buffer, format="PNG")
                    raster.append(write_binary(root, directory / f"raster-{index}-{number}.png", buffer.getvalue()))
            artifacts.append(
                {
                    "pdf": write_binary(root, directory / f"publication-{index}.pdf", pdf.read_bytes()),
                    "pages": pages,
                    "rasters": raster,
                    "pptx": pptx_bindings.get(pdf.name),
                }
            )
        return native_result(
            root,
            directory,
            capability,
            {
                "schema_version": "template-auxiliary-publication-native-v1",
                "status": "passed",
                "capability": capability,
                "artifacts": artifacts,
            },
        )


if __name__ == "__main__":
    raise SystemExit(main())

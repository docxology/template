"""Real editable projection data and fail-closed source integration regressions."""

from __future__ import annotations

import hashlib
import json
import time
from io import BytesIO
from pathlib import Path

import pytest
from pptx import Presentation

from infrastructure.rendering.pptx_deck import render_pptx
from infrastructure.rendering.slide_deck import DeckContent, Slide
from infrastructure.reporting.native_format_acceptance_contracts import NativeArtifact
from infrastructure.reporting.native_format_acceptance_checks import parse_native_bytes
from infrastructure.reporting.native_format_observer_outcome import native_observer_outcome
from infrastructure.reporting.native_pitch_source_parity import inspect_pitch_native_pairs, inspect_pitch_package
from infrastructure.reporting.native_projection_contracts import (
    compile_projection_contract,
    projection_contract_record,
    ProjectionSourceUnsupported,
)
from infrastructure.reporting.native_projection_source_dispatch import (
    observe_admitted_manuscript_projection,
    projection_control_formats,
)
from infrastructure.reporting.release_readiness_snapshot import SnapshotReader


def test_actual_package_notes_failure_reaches_enclosing_native_outcome(tmp_path: Path) -> None:
    deck = DeckContent("Measured source", slides=(Slide("Result", notes="Offset -3.25 mV."),))
    pages = compile_projection_contract(deck)
    original = render_pptx(deck, tmp_path / "original.pptx").read_bytes()
    package = Presentation(BytesIO(original))
    package.slides[1].notes_slide.notes_text_frame.text = "Offset +3.26 mV."
    stream = BytesIO()
    package.save(stream)
    changed = stream.getvalue()
    parse_native_bytes("pptx", changed)
    observed = inspect_pitch_package(changed, pages, deadline=time.monotonic() + 15, source_model=deck)
    # The deliberately unqualified native wrapper grants no execution claim.
    layout = {"status": "blocked", "findings": [], "blockers": ["actual native conversion has not run"]}
    raw = {
        "id": "pitch:long:pptx",
        "format": "pptx",
        "projection_contract": projection_contract_record(pages),
        "observation": {"conversion": {}, "raster": {}, "observation": layout, "status": "blocked"},
        "source_package": observed,
    }
    combined = native_observer_outcome(json.loads(json.dumps(raw)))
    assert combined["status"] == "fail"
    assert {row["code"] for row in combined["findings"]} == {"pitch.package.source.notes"}
    assert combined["blockers"] == layout["blockers"]
    del raw["source_package"]
    with pytest.raises(ValueError, match="must be an object"):
        native_observer_outcome(raw)


def test_all_six_actual_package_source_contracts_do_not_promote_unrun_native_pairs(tmp_path: Path) -> None:
    artifacts, contracts = {}, {}
    for length in ("short", "medium", "long"):
        deck = DeckContent(
            f"Measurement {length}", slides=(Slide("Result", bullets=("Offset -3.25 mV.",), notes=length),)
        )
        pages = compile_projection_contract(deck)
        native = render_pptx(deck, tmp_path / f"{length}.pptx").read_bytes()
        parsed = parse_native_bytes("pptx", native)
        package = inspect_pitch_package(native, pages, deadline=time.monotonic() + 15, source_model=deck)
        assert package["status"] == "pass" and parsed["slides"] == len(pages)
        for kind in ("pdf", "pptx"):
            identity = f"pitch:{length}:{kind}"
            contracts[identity] = pages
            # Only the genuine editable package is observed. Missing PDF and
            # native-converted layout remain explicit DATA blockers, not passes.
            native_observation = {
                "status": "blocked",
                "findings": [],
                "blockers": ["all-page native observer pending"],
                "pages": [],
            }
            from infrastructure.reporting.native_format_pdf_oracles import pdf_control_bytes

            actual_data = native if kind == "pptx" else pdf_control_bytes("baseline")
            parse_native_bytes(kind, actual_data)
            artifacts[identity] = {
                "id": identity,
                "format": kind,
                "input": {"sha256": hashlib.sha256(actual_data).hexdigest()},
                "projection_contract": projection_contract_record(pages),
                "observation": {"observation": native_observation} if kind == "pptx" else native_observation,
            }
            if kind == "pptx":
                artifacts[identity]["source_package"] = package
    observed = inspect_pitch_native_pairs(artifacts, contracts)
    assert observed["status"] == "fail"
    assert len(observed["pairs"]) == 3
    assert len(observed["blockers"]) >= 6
    assert {item["id"] for item in observed["findings"] if item["code"] == "pitch.pair.native.pages"} == set(artifacts)
    del artifacts["pitch:long:pptx"]
    with pytest.raises(ValueError, match="exactly all three"):
        inspect_pitch_native_pairs(artifacts, contracts)


def test_manuscript_projection_cannot_launch_without_independent_source_admission(tmp_path: Path) -> None:
    from infrastructure.reporting.native_format_pdf_oracles import pdf_control_bytes

    native = tmp_path / "output/readiness/current.pdf"
    native.parent.mkdir(parents=True)
    native.write_bytes(pdf_control_bytes("baseline"))
    reader = SnapshotReader(tmp_path)
    snapshot = reader.read(native.relative_to(tmp_path).as_posix())
    leaf = tmp_path / "output/readiness/observer"
    with pytest.raises(ProjectionSourceUnsupported, match="independent immutable source"):
        observe_admitted_manuscript_projection(
            tmp_path,
            leaf,
            "templates/template_code_project",
            NativeArtifact("section:measurement.md:slides", "pdf", snapshot.relative),
            snapshot,
            reader,
            env={},
            deadline=time.monotonic() + 15,
            run_id="unadmitted-source-negative",
            source_reference=None,
            verify_source_correspondence=None,
        )
    assert not leaf.exists()


def test_actual_pitch_source_obligations_select_projection_controls_when_manuscript_slides_are_off(
    tmp_path: Path,
) -> None:
    import yaml
    from infrastructure.reporting.native_format_acceptance_contracts import native_artifact_contract
    from tests.infra_tests.reporting.test_native_format_acceptance import _project

    scope = "templates/template_pitch_deck"
    project = _project(tmp_path, scope, slides=False)
    path = project / "manuscript/config.yaml"
    config = yaml.safe_load(path.read_text())
    config["project_config"] = {
        "deck": {
            "formats": ["pdf", "pptx"],
            "lengths": ["short", "medium", "long"],
            "pitch_subject": "template_template",
        }
    }
    path.write_text(yaml.safe_dump(config))
    selected = ("pdf", "html")
    contract = native_artifact_contract(tmp_path, scope, selected, SnapshotReader(tmp_path))
    assert len([row for row in contract if row.id.startswith("pitch:")]) == 6
    assert projection_control_formats(contract, selected) == ("pdf", "html", "pptx")
    assert not (project / "output").exists()

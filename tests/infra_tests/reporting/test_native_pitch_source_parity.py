"""Actual editable Pitch package/source controls and complete variant refusal."""

from __future__ import annotations

import time
from io import BytesIO
from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation
from pptx.util import Pt

from infrastructure.rendering.slide_deck import DeckContent, Slide
from infrastructure.rendering.pptx_deck import render_pptx
from infrastructure.reporting.native_projection_contracts import compile_projection_contract
from infrastructure.reporting.native_pitch_source_parity import inspect_pitch_package, inspect_pitch_native_pairs
from infrastructure.reporting.native_pitch_source_layout import pitch_package_layout
from infrastructure.reporting.release_readiness_snapshot import SnapshotReader


@pytest.mark.parametrize(
    "mutation,code",
    (
        ("baseline", None),
        ("number", "pitch.package.source.text"),
        ("font", "pitch.package.typography.font"),
        ("crop", "pitch.package.layout.crop"),
        ("target", "pitch.package.source.targets"),
        ("figure", "pitch.package.figure.source-pixels"),
        ("geometry", "pitch.package.figure.geometry"),
    ),
)
def test_actual_pitch_package_source_controls(tmp_path: Path, mutation: str, code: str | None) -> None:
    original = tmp_path / "output/readiness/figure.png"
    original.parent.mkdir(parents=True)
    Image.new("RGB", (128, 64), "navy").save(original)
    deck = DeckContent(
        "Measurement",
        slides=(Slide("Observed", bullets=("Offset -3.25 mV.",), figure_path=original, source="data.csv"),),
    )
    pages = compile_projection_contract(deck, source_base_url="https://example.org", reader=SnapshotReader(tmp_path))
    native = render_pptx(deck, tmp_path / "actual.pptx", source_base_url="https://example.org")
    prs = Presentation(native)
    slide = prs.slides[1]
    text = [shape for shape in slide.shapes if shape.has_text_frame]
    body = next(shape for shape in text if "Offset" in shape.text)
    if mutation == "number":
        body.text_frame.paragraphs[0].runs[0].text = "• Offset +3.26 mV."
    elif mutation == "font":
        body.text_frame.paragraphs[0].runs[0].font.size = Pt(5)
    elif mutation == "crop":
        body.left = Pt(-70)
    elif mutation == "target":
        footer = next(shape for shape in text if "Source:" in shape.text)
        footer.text_frame.paragraphs[0].runs[0].hyperlink.address = "https://example.org/wrong"
    elif mutation in {"figure", "geometry"}:
        picture = next(shape for shape in slide.shapes if shape.shape_type == 13)
        if mutation == "geometry":
            picture.width = Pt(15)
        else:
            replacement = tmp_path / "changed.png"
            Image.new("RGB", (128, 64), "red").save(replacement)
            left, top, width, height = picture.left, picture.top, picture.width, picture.height
            element = picture._element
            element.getparent().remove(element)
            slide.shapes.add_picture(str(replacement), left, top, width=width, height=height)
    stream = BytesIO()
    prs.save(stream)
    observed = inspect_pitch_package(stream.getvalue(), pages, deadline=time.monotonic() + 15)
    assert observed["blockers"] == []
    assert observed["status"] == ("pass" if code is None else "fail")
    if code:
        assert code in {item["code"] for item in observed["findings"]}


def test_pitch_pair_gate_refuses_missing_variants_before_receipt_summaries() -> None:
    with pytest.raises(ValueError, match="all three PDF/PPTX"):
        inspect_pitch_native_pairs({}, {})


def test_actual_pitch_package_source_layout_detects_in_bounds_relocation(tmp_path: Path) -> None:
    deck = DeckContent("Measurement", slides=(Slide("Observed", bullets=("Offset -3.25 mV.",)),))
    pages = compile_projection_contract(deck)
    original = render_pptx(deck, tmp_path / "source.pptx").read_bytes()
    prs = Presentation(BytesIO(original))
    body = next(shape for shape in prs.slides[1].shapes if shape.has_text_frame and "Offset" in shape.text)
    body.left += Pt(5)
    stream = BytesIO()
    prs.save(stream)
    changed = stream.getvalue()
    assert inspect_pitch_package(changed, pages, deadline=time.monotonic() + 15)["status"] == "pass"
    assert pitch_package_layout(original) != pitch_package_layout(changed)


def test_actual_pitch_package_preserves_source_speaker_notes(tmp_path: Path) -> None:
    deck = DeckContent("Measurement", slides=(Slide("Observed", notes="Offset -3.25 mV."),))
    pages = compile_projection_contract(deck)
    native = render_pptx(deck, tmp_path / "source.pptx").read_bytes()
    baseline = inspect_pitch_package(native, pages, deadline=time.monotonic() + 15, source_model=deck)
    assert baseline["status"] == "pass"
    assert baseline["source_notes_checked"] is True
    prs = Presentation(BytesIO(native))
    prs.slides[1].notes_slide.notes_text_frame.text = "Offset +3.26 mV."
    changed = BytesIO()
    prs.save(changed)
    observation = inspect_pitch_package(changed.getvalue(), pages, deadline=time.monotonic() + 15, source_model=deck)
    assert observation["status"] == "fail"
    assert "pitch.package.source.notes" in {row["code"] for row in observation["findings"]}


def test_actual_quote_package_includes_source_owned_native_quote_marks(tmp_path: Path) -> None:
    deck = DeckContent(
        "Measurement",
        slides=(Slide("", kind="quote", quote_text="Offset -3.25 mV.", quote_attribution="Measured source"),),
    )
    pages = compile_projection_contract(deck)
    native = render_pptx(deck, tmp_path / "quote.pptx").read_bytes()
    observed = inspect_pitch_package(native, pages, deadline=time.monotonic() + 15, source_model=deck)
    assert observed["status"] == "pass"
    assert "“" in observed["pages"][1]["text_tokens"]
    assert "—" in observed["pages"][1]["text_tokens"]

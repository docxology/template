"""Genuine source-model and native PDF projection controls without replacement mocks."""

from __future__ import annotations

import time
from io import BytesIO
from pathlib import Path

import pytest
from reportlab.pdfgen.canvas import Canvas

from infrastructure.rendering.slide_deck import DeckContent, Slide
from infrastructure.reporting.native_projection_contracts import compile_projection_contract
from infrastructure.reporting.native_projection_checks import inspect_projection


def test_projection_preserves_order_duplicates_and_targets() -> None:
    deck = DeckContent(
        "Invented",
        slides=(
            Slide("Value", bullets=("-3.25 mV", "-3.25 mV"), source="data.csv", qr_url="https://example.org/slide"),
        ),
    )
    pages = compile_projection_contract(deck, source_base_url="https://example.org")
    assert len(pages) == 2
    assert pages[1].texts[1:3] == (("body", "• -3.25 mV", 18.0), ("body", "• -3.25 mV", 18.0))
    assert pages[1].citation == "https://example.org/data.csv"
    assert pages[1].qr_target == "https://example.org/slide"


def test_projection_refuses_incomplete_native_pages() -> None:
    stream = BytesIO()
    canvas = Canvas(stream)
    canvas.drawString(20, 40, "Invented")
    canvas.save()
    pages = compile_projection_contract(DeckContent("Invented", slides=(Slide("Actual"),)))
    with pytest.raises(ValueError, match="complete source/page/raster"):
        inspect_projection(stream.getvalue(), (), pages, deadline=time.monotonic() + 10)


def test_native_office_unavailable_is_not_structural_success(tmp_path: Path) -> None:
    from infrastructure.reporting.native_format_acceptance_checks import NativeToolUnavailable, native_tool_record

    with pytest.raises(NativeToolUnavailable, match="not installed"):
        native_tool_record(
            tmp_path,
            tmp_path,
            "soffice",
            ["--version"],
            env={"PATH": ""},
            deadline=time.monotonic() + 10,
            run_id="invented",
            purpose="version",
        )


@pytest.mark.parametrize("dimensions", ((200, 80), (80, 200)))
def test_source_figure_geometry_matches_native_pdf_and_pptx(tmp_path: Path, dimensions: tuple[int, int]) -> None:
    from PIL import Image
    from pptx import Presentation
    from infrastructure.rendering.slide_deck import render_pdf
    from infrastructure.rendering.pptx_deck import render_pptx
    from infrastructure.reporting.release_readiness_snapshot import SnapshotReader
    from infrastructure.reporting.native_projection_checks import inspect_embedded_figure

    source = tmp_path / "output/readiness/figure.png"
    source.parent.mkdir(parents=True)
    Image.new("RGB", dimensions, "navy").save(source)
    deck = DeckContent("Invented", slides=(Slide("Figure", figure_path=source),))
    contract = compile_projection_contract(deck, reader=SnapshotReader(tmp_path))[-1].figure_projection
    assert contract is not None
    native_pdf = render_pdf(deck, tmp_path / "figure.pdf")
    assert inspect_embedded_figure(native_pdf.read_bytes(), 2, contract)["findings"] == []
    # Inventory-only nested resources must not be mistaken for a second draw.
    # Real native readers can retain the same image through several form trees.
    from pypdf import PdfReader, PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject, ArrayObject, NumberObject

    parsed = PdfReader(native_pdf)
    writer = PdfWriter()
    for page in parsed.pages:
        writer.add_page(page)
    page = writer.pages[1]
    references = page["/Resources"]["/XObject"]
    form = DecodedStreamObject()
    form.set_data(b"")
    form.update(
        {
            NameObject("/Type"): NameObject("/XObject"),
            NameObject("/Subtype"): NameObject("/Form"),
            NameObject("/BBox"): ArrayObject([NumberObject(0), NumberObject(0), NumberObject(720), NumberObject(405)]),
            NameObject("/Resources"): DictionaryObject({NameObject("/XObject"): DictionaryObject(dict(references))}),
        }
    )
    references[NameObject("/UnusedInventory")] = writer._add_object(form)
    repeated = BytesIO()
    writer.write(repeated)
    assert inspect_embedded_figure(repeated.getvalue(), 2, contract)["findings"] == []
    # Drawing the extra occurrence is a real geometry failure, even though its
    # underlying image bytes and repeated resource inventory are unchanged.
    image_name = next(iter(form["/Resources"]["/XObject"]))
    form.set_data(f"{image_name} Do".encode())
    extra = DecodedStreamObject()
    extra.set_data(b"q /UnusedInventory Do Q")
    page[NameObject("/Contents")] = ArrayObject([page.raw_get("/Contents"), writer._add_object(extra)])
    duplicated = BytesIO()
    writer.write(duplicated)
    assert {finding["code"] for finding in inspect_embedded_figure(duplicated.getvalue(), 2, contract)["findings"]} == {
        "projection.figure.geometry"
    }
    native_pptx = render_pptx(deck, tmp_path / "figure.pptx")
    picture = next(shape for shape in Presentation(native_pptx).slides[1].shapes if shape.shape_type == 13)
    left, top, right, bottom = contract.box_pt
    assert picture.left / 12700 == pytest.approx(left, abs=0.001)
    assert picture.top / 12700 == pytest.approx(top, abs=0.001)
    assert picture.width / 12700 == pytest.approx(right - left, abs=0.001)
    assert picture.height / 12700 == pytest.approx(bottom - top, abs=0.001)


@pytest.mark.parametrize(
    "name",
    (
        "baseline",
        "changed-number",
        "clipped-text",
        "tiny-text",
        "opaque-overlay",
        "wrong-citation",
        "changed-figure",
        "changed-figure-detail",
        "distorted-figure",
    ),
)
def test_actual_valid_projection_controls(tmp_path: Path, name: str) -> None:
    import os
    import shutil
    from infrastructure.reporting.native_projection_checks import run_projection_control

    if not shutil.which("pdftoppm") or not shutil.which("pdffonts"):
        pytest.skip("actual native PDF tools are unavailable")
    directory = tmp_path / "output/readiness/control"
    directory.mkdir(parents=True)
    result = run_projection_control(
        tmp_path, directory, name, env=dict(os.environ), deadline=time.monotonic() + 30, run_id="invented-control"
    )
    assert result["parser"]["pages"] == 1
    assert len(result["raster"]["rasters"]) == 1
    assert result["oracle_exit_code"] == (0 if name == "baseline" else 1)
    assert result["observation"]["status"] == ("pass" if name == "baseline" else "fail")


@pytest.mark.timeout(90)
@pytest.mark.parametrize(
    "name",
    (
        "baseline",
        "changed-number",
        "tiny-text",
        "clipped-text",
        "opaque-overlay",
        "wrong-citation",
        "changed-figure",
        "changed-figure-detail",
        "distorted-figure",
    ),
)
def test_supervised_actual_projection_controls(name: str) -> None:
    import os
    import shutil
    import sys
    import tempfile
    from infrastructure.reporting.native_projection_checks import supervise_projection_controls

    if not shutil.which("pdftoppm") or not shutil.which("pdffonts"):
        pytest.skip("actual native PDF tools are unavailable")
    root = Path(__file__).absolute().parents[3]
    parent = root / "output/readiness"
    parent.mkdir(parents=True, exist_ok=True)
    # Preserve actual failed child receipts for diagnosis. Successful test
    # evidence is disposable; an original timeout or setup failure is not.
    directory = Path(tempfile.mkdtemp(prefix="projection-control-test-", dir=parent))
    records = supervise_projection_controls(
        root,
        directory,
        interpreter=Path(sys.executable),
        env=dict(os.environ),
        deadline=time.monotonic() + 60,
        run_id="invented-supervised-control",
        include_qr=False,
        control_names=(name,),
    )
    assert len(records) == 1
    assert records[0]["name"] == name
    assert records[0]["original_exit_code"] == (0 if name == "baseline" else 1)
    shutil.rmtree(directory)


@pytest.mark.parametrize(
    "name",
    ("baseline", "changed-number", "reordered-spine", "omitted-text", "tiny-text", "clipped-text", "opaque-overlay"),
)
def test_epub_valid_spine_controls_preserve_package_validity(name: str) -> None:
    from infrastructure.reporting.native_epub_observer import epub_control_bytes, inspect_epub_spine
    from infrastructure.reporting.native_format_acceptance_checks import parse_native_bytes

    content, chapters = epub_control_bytes(name)
    repeated, repeated_chapters = epub_control_bytes(name)
    assert content == repeated and chapters == repeated_chapters
    parsed = parse_native_bytes("epub", content)
    assert parsed["xhtml_documents"] == 3
    observed = inspect_epub_spine(content, chapters)
    assert len(observed["spine"]) == 2
    if name in {"changed-number", "reordered-spine", "omitted-text"}:
        assert observed["status"] == "fail"
        assert {finding["code"] for finding in observed["findings"]} == {"epub.spine.content"}
    else:
        assert observed["status"] == "pass"
    # CSS mutants remain semantic matches: only an actual reader's glyph/pixel
    # observations may establish their intended layout rejection.
